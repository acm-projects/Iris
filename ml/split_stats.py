"""Count train / val / test videos per word for each WLASL subset and save a spreadsheet.

Reads data/splits/asl100.json, asl300.json, asl1000.json and asl2000.json and
writes one Excel workbook with a tab per subset (one row per word) plus a
Summary tab comparing the four subsets.

  python ml/split_stats.py
  python ml/split_stats.py --output ml/wlasl_split_counts.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Cue repo root (this file lives in Cue/ml/), used for default paths.
REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
SUBSETS = ["asl100", "asl300", "asl1000", "asl2000"]
SPLITS = ["train", "val", "test"]

# Styling: Arial everywhere, a dark header row, bold totals row.
FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="305496")
TOTAL_FILL = PatternFill("solid", start_color="D9E1F2")
TOP_BORDER = Border(top=Side(style="thin"))


def count_splits(split_file: str, pose_root: str | None) -> list[dict]:
    """Return one row per gloss: {gloss, train, val, test, missing_pose}, in the file's order."""
    with open(split_file) as f:
        content = json.load(f)

    rows = []
    for entry in content:
        counts = Counter(inst["split"] for inst in entry["instances"])
        # How many of this word's videos have no pose folder on disk (0 if we
        # can't check because the pose data isn't present).
        missing = 0
        if pose_root:
            missing = sum(not os.path.isdir(os.path.join(pose_root, inst["video_id"]))
                          for inst in entry["instances"])
        rows.append({"gloss": entry["gloss"], **{s: counts.get(s, 0) for s in SPLITS}, "missing_pose": missing})
    return rows


def style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"  # header stays visible while scrolling


def write_subset_sheet(wb: Workbook, name: str, rows: list[dict]) -> None:
    """One tab per subset: Rank, Gloss, Train, Val, Test, Total, Train %, Val %, Test %, Missing pose."""
    ws = wb.create_sheet(name)
    headers = ["Rank", "Gloss", "Train", "Val", "Test", "Total", "Train %", "Val %", "Test %", "Missing pose files"]
    ws.append(headers)
    style_header(ws, len(headers))

    # Most videos first; ties broken alphabetically so the order is stable.
    rows = sorted(rows, key=lambda r: (-(r["train"] + r["val"] + r["test"]), r["gloss"]))

    for i, r in enumerate(rows, start=1):
        row = i + 1  # sheet row (row 1 is the header)
        ws.append([
            i, r["gloss"], r["train"], r["val"], r["test"],
            f"=SUM(C{row}:E{row})",                       # Total
            f"=IF(F{row}=0,0,C{row}/F{row})",             # Train %
            f"=IF(F{row}=0,0,D{row}/F{row})",             # Val %
            f"=IF(F{row}=0,0,E{row}/F{row})",             # Test %
            r["missing_pose"],
        ])

    # Totals row under the table.
    last = len(rows) + 1
    total_row = last + 1
    ws.append([
        None, "TOTAL",
        f"=SUM(C2:C{last})", f"=SUM(D2:D{last})", f"=SUM(E2:E{last})", f"=SUM(F2:F{last})",
        f"=IF(F{total_row}=0,0,C{total_row}/F{total_row})",
        f"=IF(F{total_row}=0,0,D{total_row}/F{total_row})",
        f"=IF(F{total_row}=0,0,E{total_row}/F{total_row})",
        f"=SUM(J2:J{last})",
    ])

    # Fonts and number formats.
    for r_cells in ws.iter_rows(min_row=2, max_row=total_row):
        for cell in r_cells:
            cell.font = FONT
            if cell.column in (7, 8, 9):
                cell.number_format = "0.0%"
            elif cell.column >= 3:
                cell.number_format = "#,##0"
    for cell in ws[total_row]:
        cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, TOP_BORDER

    widths = [7, 18, 9, 9, 9, 9, 10, 10, 10, 13]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"  # sort/filter dropdowns


def write_summary_sheet(wb: Workbook, sheet_rows: dict[str, int]) -> None:
    """Compare the subsets. Every number is a formula that reads the subset tabs."""
    ws = wb.create_sheet("Summary", 0)
    headers = ["Subset", "Words", "Train videos", "Val videos", "Test videos", "Total videos",
               "Min videos / word", "Median videos / word", "Max videos / word", "Words with < 10 train"]
    ws.append(headers)
    style_header(ws, len(headers))

    for i, name in enumerate(SUBSETS, start=2):
        last = sheet_rows[name] + 1  # last data row on that tab
        total = last + 1             # its TOTAL row
        s = f"'{name}'"
        ws.append([
            name,
            f"=COUNTA({s}!B2:B{last})",
            f"={s}!C{total}", f"={s}!D{total}", f"={s}!E{total}", f"={s}!F{total}",
            f"=MIN({s}!F2:F{last})", f"=MEDIAN({s}!F2:F{last})", f"=MAX({s}!F2:F{last})",
            f'=COUNTIF({s}!C2:C{last},"<10")',
        ])
        for cell in ws[i]:
            cell.font = FONT
            if cell.column >= 2:
                cell.number_format = "#,##0" if cell.column != 8 else "#,##0.0"

    note_row = len(SUBSETS) + 3
    ws.cell(row=note_row, column=1,
            value="Counts come from the WLASL split files in data/splits/ "
                  "(https://github.com/dxli94/WLASL). Each subset contains the smaller ones: "
                  "asl2000 includes every asl100 video with the same split label.").font = Font(name="Arial", size=9, italic=True)

    widths = [10, 8, 13, 11, 11, 13, 15, 18, 15, 18]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data-root", default=os.path.join(REPO_ROOT, "data"),
                    help="folder containing splits/ (and optionally pose_per_individual_videos/)")
    ap.add_argument("--output", default=os.path.join(REPO_ROOT, "ml", "wlasl_split_counts.xlsx"))
    args = ap.parse_args()

    pose_root = os.path.join(args.data_root, "pose_per_individual_videos")
    if not os.path.isdir(pose_root):
        print(f"Note: {pose_root} not found, so 'Missing pose files' will be 0.")
        pose_root = None

    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    sheet_rows = {}
    for name in SUBSETS:
        rows = count_splits(os.path.join(args.data_root, "splits", f"{name}.json"), pose_root)
        write_subset_sheet(wb, name, rows)
        sheet_rows[name] = len(rows)
        totals = {s: sum(r[s] for r in rows) for s in SPLITS}
        print(f"{name}: {len(rows)} words, train {totals['train']}, val {totals['val']}, test {totals['test']}")

    write_summary_sheet(wb, sheet_rows)
    # Cells hold formulas without cached results, so ask Excel / Numbers /
    # Sheets to compute everything when the file is opened.
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.output)
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
